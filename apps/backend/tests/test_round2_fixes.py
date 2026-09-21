"""Regression tests for the Sept-2026 audit round (offline: no Redis / DB / LLM needed)."""
import asyncio
import io
import json
import os

import fakeredis
import openpyxl
import pandas as pd
import pytest

from agents.cleaner import clean_records
from agents.validator import score_dataset
from ai.providers.groq_client import parse_rate_limit_error
from api.upload import validate_file_magic_bytes
from core.dataset_query import apply_query, dataset_stats
from core.export_utils import build_frame, render_export, sanitize_for_export
from core.llm_json import extract_json, extract_records
from memory.session_store import RedisStore
from parsers.document_parser import DocumentParseError, parse_tabular_records, parse_upload


@pytest.fixture
def store():
    async def _make():
        s = RedisStore()
        s.redis_client = fakeredis.FakeAsyncRedis(decode_responses=True)
        s._loop = asyncio.get_running_loop()
        return s
    return _make


# ── upload / parsing ─────────────────────────────────────────────────────────
def test_utf8_multibyte_boundary_is_not_rejected():
    telugu = ("పేరు,ధర,వివరణ\n" * 60).encode("utf-8")
    assert validate_file_magic_bytes(telugu[:512], ".csv") is True
    assert validate_file_magic_bytes(b"name,price\x00\x00\xff\xfe", ".csv") is False


def test_xls_variants(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook(); ws = wb.add_sheet("S"); ws.write(0, 0, "Name"); ws.write(1, 0, "W")
    real = tmp_path / "real.xls"; wb.save(str(real))
    assert parse_tabular_records(str(real)) == [{"Name": "W"}]
    html = tmp_path / "html.xls"
    html.write_text("<html><table><tr><th>N</th></tr><tr><td>a</td></tr></table></html>")
    assert parse_tabular_records(str(html)) == [{"N": "a"}]
    tsv = tmp_path / "tsv.xls"; tsv.write_bytes("\ufeffN\tP\nq\t2\n".encode("utf-8"))
    assert parse_tabular_records(str(tsv)) == [{"N": "q", "P": 2}]


def test_csv_encodings_and_leading_zeros(tmp_path):
    p = tmp_path / "a.csv"; p.write_bytes("పేరు,ID\nరవి,007\n".encode("utf-8"))
    assert parse_tabular_records(str(p)) == [{"పేరు": "రవి", "ID": "007"}]
    q = tmp_path / "b.csv"; q.write_bytes("Name;Price\nCafé;2\n".encode("cp1252"))
    assert parse_tabular_records(str(q)) == [{"Name": "Café", "Price": 2}]
    text, rows = parse_upload(str(p)); assert rows and "007" in text


def test_empty_file_raises_friendly_error(tmp_path):
    p = tmp_path / "e.csv"; p.write_bytes(b"a,b\n")
    with pytest.raises(DocumentParseError):
        parse_tabular_records(str(p))


# ── export ───────────────────────────────────────────────────────────────────
DATA = [{"name": "రవి", "price": -5, "tags": ["a", "b"], "specs": {"ram": "8GB"}, "f": "=1+1"}, {"name": "B", "extra": "x"}]


def test_export_csv_has_bom_union_columns_and_readable_lists():
    raw = render_export("csv", DATA)
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert text.splitlines()[0] == "name,price,tags,specs,f,extra"
    assert "a; b" in text and "['a'" not in text and "'=1+1" in text


def test_export_xlsx_is_real_workbook_with_text_formulas():
    wb = openpyxl.load_workbook(io.BytesIO(render_export("excel", DATA)))
    ws = wb.active
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold
    cell = [c for c in ws[2] if c.value == "=1+1"][0]
    assert cell.data_type == "s"


def test_export_json_is_not_mangled():
    rows = json.loads(render_export("json", [{"a": "-5", "b": 5}]))
    assert rows == [{"a": "-5", "b": 5}]


def test_sanitize_keeps_numbers_neutralises_formulas():
    assert [sanitize_for_export(v) for v in ["-5", "+1,200.50", "-", "=SUM(1)", "@x", "+cmd|x"]] == \
           ["-5", "+1,200.50", "-", "'=SUM(1)", "'@x", "'+cmd|x"]


# ── LLM plumbing ─────────────────────────────────────────────────────────────
def test_rate_limit_parser():
    class E(Exception):
        message = "on tokens per minute (TPM): Limit 8000. Please try again in 7.5s."
    assert parse_rate_limit_error(E()) == (8, "minute")
    class D(Exception):
        message = "on tokens per day (TPD): Limit 200000. Please try again in 12m30.2s."
    assert parse_rate_limit_error(D()) == (750, "day")


def test_llm_json_helpers():
    assert extract_json('```json\n{"a":1}\n```') == {"a": 1}
    assert extract_records('{"records":[{"a":1},{"a":2}]}') == [{"a": 1}, {"a": 2}]
    assert extract_records('{"records":[{"a":1},{"b":2},{"c":') == [{"a": 1}, {"b": 2}]   # truncated output salvaged


def test_dataset_query_runs_on_all_rows():
    rows = [{"n": "A", "price": "₹1,200"}, {"n": "B", "price": "₹300"}, {"n": "C", "price": "₹4,999"}]
    out = apply_query(rows, {"filters": [{"column": "Price", "op": "<=", "value": 1500}], "sort": [{"column": "price", "order": "asc"}]})
    assert [r["n"] for r in out] == ["B", "A"]
    assert dataset_stats(rows)["price"]["sum"] == 6499.0


def test_cleaner_and_validator_are_deterministic():
    rows = clean_records([{"title": " A ", "price": "₹1,299", "url": "/p"}, {"title": "A", "price": "₹1,299", "url": "/p"}], "https://s.com")
    assert rows == [{"title": "A", "price": 1299, "currency": "INR", "url": "https://s.com/p"}]
    assert 0 <= score_dataset(rows, ["title", "price"])["confidence_score"] <= 100


# ── Redis ────────────────────────────────────────────────────────────────────
@pytest.mark.asyncio
async def test_dataset_roundtrip_claim_and_titles(store):
    s = await store()
    await s.save_session_data("s1", {"cleaned_data": [{"a": "రవి"}]})
    assert (await s.get_session_data("s1"))["cleaned_data"] == [{"a": "రవి"}]
    assert await s.claim_session("c", "u1") and not await s.claim_session("c", "u2")
    await s.add_user_session("u1", "c"); await s.set_session_title_if_missing("c", "My first chat")
    assert (await s.get_user_sessions("u1"))[0]["title"] == "My first chat"
    await s.delete_session("c")
    assert await s.get_user_sessions("u1") == []


# ── Performance-related behaviour ────────────────────────────────────────────
@pytest.mark.asyncio
async def test_fast_planner_skips_the_llm(monkeypatch):
    from unittest.mock import AsyncMock
    from agents.planner import planner_agent
    from core.config import settings
    monkeypatch.setattr(settings, "PLANNER_MODE", "fast")
    boom = AsyncMock(side_effect=AssertionError("planner LLM must not be called in fast mode"))
    monkeypatch.setattr("agents.planner.ai_router.generate", boom)
    out = await planner_agent.run({"user_request": "get prices", "target_url": "https://example.com", "metadata": {}}, "s")
    assert out["extraction_goal"] == "get prices" and out["target_url"] == "https://example.com"
    boom.assert_not_called()


def test_gzip_and_timing_headers_are_enabled():
    import httpx, asyncio
    import main

    async def go():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://t") as c:
            return await c.get("/health", headers={"Accept-Encoding": "gzip"})
    r = asyncio.run(go())
    assert r.status_code == 200
    assert "server-timing" in {k.lower() for k in r.headers}
    assert any(type(m.cls).__name__ == "type" and m.cls.__name__ == "GZipMiddleware" for m in main.app.user_middleware)
