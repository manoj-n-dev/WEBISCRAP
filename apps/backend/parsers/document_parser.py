import csv
import datetime
import io
import math
import os
import re
from typing import Any, Dict, List, Optional, Tuple, cast

import pandas as pd
from loguru import logger

# Heavy / optional imports are done lazily so a missing OCR stack never breaks CSV/XLSX support.
MAX_IMAGE_PIXELS = 50_000_000          # ~50 MP: covers modern phone cameras (12-48 MP); bomb guard at 2x
MAX_PDF_PAGES = 50
MAX_SPREADSHEET_ROWS = 5000
MAX_SHEETS = 5
MAX_DOCX_PARAGRAPHS = 2000
MAX_OCR_SIDE = 3000

TABULAR_EXTENSIONS = (".csv", ".xlsx", ".xls")


class DocumentParseError(Exception):
    """Raised with a user-safe message when a file cannot be parsed."""


# ─── Tabular files (CSV / XLSX / XLS) ─────────────────────────────────────────

def _decode_bytes(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample.splitlines() else ""
        return max((",", ";", "\t", "|"), key=first.count)


_NUM_RE = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d+)?$")


def _infer_numeric_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Convert string columns to numbers only when EVERY value is a clean number
    (keeps IDs / phone numbers with leading zeros as text)."""
    for col in df.columns:
        s = df[col]
        if s.dtype != object:
            continue
        vals = s.dropna().astype(str).str.strip()
        vals = vals[vals != ""]
        if len(vals) and vals.map(lambda v: bool(_NUM_RE.match(v))).all():
            num = pd.to_numeric(s.astype(str).str.strip().replace({"": None}), errors="coerce")
            df[col] = num.astype("Int64").astype(object).where(num.notna(), cast(Any, None)) if (num.dropna() % 1 == 0).all() else num
    return df


def _read_text_table(raw: bytes) -> pd.DataFrame:
    text = _decode_bytes(raw)
    stripped = text.lstrip()[:200].lower()
    if stripped.startswith("<") and ("<table" in text.lower() or "<html" in stripped):
        tables = pd.read_html(io.StringIO(text))
        if not tables:
            raise DocumentParseError("No table found in this file.")
        return tables[0].head(MAX_SPREADSHEET_ROWS)
    delim = _sniff_delimiter(text)
    df = pd.read_csv(io.StringIO(text), sep=delim, nrows=MAX_SPREADSHEET_ROWS, dtype=str,
                     keep_default_na=False, na_values=[""], engine="python", on_bad_lines="skip")
    return _infer_numeric_columns(df)


def _read_excel_frames(path: str, ext: str, raw_head: bytes) -> List[Tuple[str, pd.DataFrame]]:
    if raw_head.startswith(b"PK"):
        engine = "openpyxl"
    elif raw_head.startswith(b"\xd0\xcf\x11\xe0"):
        engine = "xlrd"
    else:
        with open(path, "rb") as f:                  # HTML / XML / CSV saved with an .xls extension
            return [("Sheet1", _read_text_table(f.read()))]
    try:
        with pd.ExcelFile(path, engine=engine) as xl:
            frames = []
            for name in xl.sheet_names[:MAX_SHEETS]:
                df = xl.parse(name, nrows=MAX_SPREADSHEET_ROWS)
                if not df.dropna(how="all").empty:
                    frames.append((str(name), df))
    except ImportError as e:
        raise DocumentParseError("Legacy .xls support is not installed on the server (missing 'xlrd').") from e
    if not frames:
        raise DocumentParseError("The spreadsheet is empty.")
    return frames


def read_tabular_file(file_path: str) -> pd.DataFrame:
    """Load CSV / XLSX / XLS into ONE DataFrame (multi-sheet workbooks get a `sheet` column)."""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        with open(file_path, "rb") as f:
            head = f.read(16)
        if ext == ".csv":
            with open(file_path, "rb") as f:
                df = _read_text_table(f.read())
        else:
            frames = _read_excel_frames(file_path, ext, head)
            if len(frames) == 1:
                df = frames[0][1]
            else:
                parts = []
                for name, part in frames:
                    part = part.copy()
                    part.insert(0, "sheet", name)
                    parts.append(part)
                df = pd.concat(parts, ignore_index=True)
    except DocumentParseError:
        raise
    except Exception as e:
        logger.error(f"Tabular parse failed for {file_path}: {e}")
        raise DocumentParseError(f"Could not read this {ext.lstrip('.').upper()} file. Check that it is a valid, non-password-protected file.") from e

    df = df.dropna(how="all").dropna(axis=1, how="all").head(MAX_SPREADSHEET_ROWS)
    if df.empty:
        raise DocumentParseError("No rows were found in this file.")
    seen: Dict[str, int] = {}
    names = []
    for i, c in enumerate(df.columns):
        name = str(c).strip()
        if not name or name.lower().startswith("unnamed:"):
            name = f"column_{i + 1}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    df.columns = names
    return df.reset_index(drop=True)


def _json_safe(v: Any) -> Any:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, (pd.Timestamp, datetime.datetime, datetime.date)):
        return v.isoformat()
    if hasattr(v, "item"):
        v = v.item()
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


def dataframe_to_records(df: pd.DataFrame) -> List[Dict[str, Any]]:
    cols = list(df.columns)
    return [{c: _json_safe(v) for c, v in zip(cols, row)} for row in df.itertuples(index=False, name=None)]


def parse_tabular_records(file_path: str) -> List[Dict[str, Any]]:
    return dataframe_to_records(read_tabular_file(file_path))


def _df_to_text(df: pd.DataFrame) -> str:
    return df.to_csv(index=False)


def parse_csv(file_path: str) -> str:
    """CSV -> compact CSV text (no fixed-width padding; ~3x fewer tokens than DataFrame.to_string)."""
    return _df_to_text(read_tabular_file(file_path))


def parse_excel(file_path: str) -> str:
    return _df_to_text(read_tabular_file(file_path))


# ─── Documents / images ───────────────────────────────────────────────────────

def parse_pdf(file_path: str) -> str:
    import pdfplumber
    text_content = []
    try:
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            for page in pdf.pages[:MAX_PDF_PAGES]:
                text = page.extract_text()
                if text:
                    text_content.append(text)
            if total_pages > MAX_PDF_PAGES:
                text_content.append(f"\n[Notice: Truncated at {MAX_PDF_PAGES} of {total_pages} total pages.]")
        result = "\n".join(text_content).strip()
        if not result:
            raise DocumentParseError("No selectable text was found in this PDF (it may be a scan). Upload it as an image instead.")
        return result
    except DocumentParseError:
        raise
    except Exception as e:
        logger.error(f"Error parsing PDF {file_path}: {e}")
        raise DocumentParseError("Could not read this PDF. It may be corrupted or password-protected.") from e


def parse_docx(file_path: str) -> str:
    import docx
    try:
        doc = docx.Document(file_path)
        paragraphs = doc.paragraphs[:MAX_DOCX_PARAGRAPHS]
        text = "\n".join(p.text for p in paragraphs if p.text.strip())
        for table in doc.tables[:20]:
            for row in table.rows:
                text += "\n" + " | ".join(c.text.strip() for c in row.cells)
        if len(doc.paragraphs) > MAX_DOCX_PARAGRAPHS:
            text += f"\n[Notice: Truncated at {MAX_DOCX_PARAGRAPHS} paragraphs.]"
        if not text.strip():
            raise DocumentParseError("No text was found in this document.")
        return text
    except DocumentParseError:
        raise
    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path}: {e}")
        raise DocumentParseError("Could not read this Word document.") from e


def parse_image(file_path: str) -> str:
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        import pytesseract
        with Image.open(file_path) as img:
            if img.format == "JPEG":
                img.draft("RGB", (MAX_OCR_SIDE, MAX_OCR_SIDE))     # cheap DCT-domain downscale for big photos
            img = ImageOps.exif_transpose(img)
            img.thumbnail((MAX_OCR_SIDE, MAX_OCR_SIDE))
            text = pytesseract.image_to_string(ImageOps.grayscale(img))
        if not text.strip():
            raise DocumentParseError("No readable text was found in this image.")
        return text
    except DocumentParseError:
        raise
    except Exception as e:
        logger.error(f"Error parsing image {file_path}: {e}")
        raise DocumentParseError("Could not read text from this image.") from e


def extract_text_from_file(file_path: str) -> str:
    """Main routing function: any supported file -> text for the LLM."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    if ext == ".docx":
        return parse_docx(file_path)
    if ext == ".csv":
        return parse_csv(file_path)
    if ext in (".xlsx", ".xls"):
        return parse_excel(file_path)
    if ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp"):
        return parse_image(file_path)
    if ext in (".txt", ".md", ".json"):
        with open(file_path, "rb") as f:
            return _decode_bytes(f.read())
    raise DocumentParseError(f"Unsupported file type: {ext}")


def parse_upload(file_path: str) -> Tuple[str, Optional[List[Dict[str, Any]]]]:
    """Upload entry-point. Returns (text_for_llm, records_or_None).
    Tabular files return exact rows so they never need an (lossy, token-hungry) LLM extraction."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext in TABULAR_EXTENSIONS:
        df = read_tabular_file(file_path)
        return _df_to_text(df), dataframe_to_records(df)
    return extract_text_from_file(file_path), None
